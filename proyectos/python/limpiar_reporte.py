"""Limpia una exportación de ventas desordenada (CSV o Excel) y genera un reporte en Excel.

Uso:
    python limpiar_reporte.py ventas_sucias.csv --salida reporte_limpio.xlsx

Qué hace:
  - Normaliza nombres de columna (acentos, mayúsculas, sinónimos como "importe" -> "monto").
  - Convierte fechas en formatos mezclados (03/01/2026, 2026-01-03, "3 ene 2026", "Jan 3, 2026").
  - Convierte montos escritos como texto ("$1,234.50", "1.234,50", "(250)", "1200 MXN").
  - Limpia texto (espacios dobles, mayúsculas inconsistentes) y unifica variantes del mismo
    nombre ("papeleria sol" y "Papelería Sol" quedan como un solo cliente).
  - Elimina duplicados exactos y separa las filas con datos faltantes o inválidos
    en una hoja "Problemas" en vez de borrarlas sin avisar.
  - Escribe un Excel con: Datos limpios, Resumen mensual, Resumen por cliente,
    Resumen por producto y Problemas.
Requiere: pandas, openpyxl.
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import pandas as pd

# Sinónimos aceptados para cada columna estándar (comparación sin acentos ni mayúsculas).
SINONIMOS = {
    "fecha": ["fecha", "date", "fecha venta", "fecha de venta", "dia"],
    "cliente": ["cliente", "customer", "nombre cliente", "client"],
    "producto": ["producto", "product", "articulo", "item", "descripcion"],
    "cantidad": ["cantidad", "qty", "quantity", "unidades", "cant"],
    "monto": ["monto", "importe", "total", "amount", "venta", "precio total"],
}
REQUERIDAS = ["fecha", "cliente", "producto", "monto"]
PARTICULAS = {"de", "del", "la", "las", "los", "el", "y", "e"}   # van en minúscula dentro de nombres

MESES = {
    "ene": 1, "jan": 1, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "aug": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12, "dec": 12,
}


def sin_acentos(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c))


def normalizar_columnas(df: pd.DataFrame) -> pd.DataFrame:
    mapa = {}
    for col in df.columns:
        clave = re.sub(r"\s+", " ", sin_acentos(str(col)).strip().lower().replace("_", " "))
        for estandar, sinonimos in SINONIMOS.items():
            if clave in sinonimos and estandar not in mapa.values():
                mapa[col] = estandar
                break
    return df.rename(columns=mapa)


def parsear_fecha(valor, dia_primero: bool = True) -> date | None:
    """Devuelve una fecha o None. Asume día/mes/año para fechas ambiguas (convención en México)."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = sin_acentos(str(valor)).strip().lower().replace(".", "")
    if not texto:
        return None
    # 2026-01-03 o 2026/01/03
    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", texto)
    if m:
        a, mes, d = map(int, m.groups())
        return _fecha_valida(a, mes, d)
    # 03/01/2026, 3-1-26
    m = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})", texto)
    if m:
        x, y, a = map(int, m.groups())
        a = a + 2000 if a < 100 else a
        d, mes = (x, y) if dia_primero else (y, x)
        return _fecha_valida(a, mes, d)
    # "3 ene 2026", "3 de enero de 2026"
    m = re.fullmatch(r"(\d{1,2})\s+(?:de\s+)?([a-z]+)\s+(?:de\s+)?(\d{4})", texto)
    if m and m.group(2)[:3] in MESES:
        return _fecha_valida(int(m.group(3)), MESES[m.group(2)[:3]], int(m.group(1)))
    # "jan 3, 2026", "january 3 2026"
    m = re.fullmatch(r"([a-z]+)\s+(\d{1,2}),?\s+(\d{4})", texto)
    if m and m.group(1)[:3] in MESES:
        return _fecha_valida(int(m.group(3)), MESES[m.group(1)[:3]], int(m.group(2)))
    return None


def _fecha_valida(a: int, mes: int, d: int) -> date | None:
    try:
        return date(a, mes, d)
    except ValueError:
        return None


def parsear_monto(valor) -> float | None:
    """Convierte '$1,234.50', '1.234,50', '(250)', '1200 MXN' a número. None si no se puede."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    negativo = (texto.startswith("(") and texto.endswith(")")) or texto.startswith("-")
    texto = re.sub(r"[^\d.,]", "", texto)
    if not texto or not re.search(r"\d", texto):
        return None
    if "," in texto and "." in texto:
        # el último separador es el decimal
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        # 1,234 o 1,234,567 -> miles; 12,5 -> decimal
        texto = texto.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", texto) else texto.replace(",", ".")
    elif texto.count(".") > 1:
        texto = texto.replace(".", "")
    try:
        numero = float(texto)
    except ValueError:
        return None
    return -numero if negativo else numero


def limpiar_texto(valor, titulo: bool = False) -> str | None:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    if not texto:
        return None
    if not titulo:
        return texto[0].upper() + texto[1:].lower()
    palabras = texto.lower().split(" ")
    return " ".join(p if i and p in PARTICULAS else p.capitalize() for i, p in enumerate(palabras))


def unificar_variantes(serie: pd.Series) -> tuple[pd.Series, int]:
    """Une variantes del mismo nombre que solo difieren en acentos, mayúsculas o espacios.
    Prefiere la escritura con acentos (suele ser la correcta en español) y luego la más frecuente."""
    clave = serie.map(lambda v: sin_acentos(v).lower() if isinstance(v, str) else None)
    canon = {}
    for k, grupo in serie.groupby(clave):
        conteo = grupo.value_counts()
        canon[k] = max(conteo.index, key=lambda v: (v != sin_acentos(v), conteo[v]))
    nueva = serie.map(lambda v: canon[sin_acentos(v).lower()] if isinstance(v, str) else v)
    return nueva, int((nueva != serie).sum())


def limpiar(df: pd.DataFrame, dia_primero: bool = True) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Devuelve (datos_limpios, problemas, estadisticas)."""
    df = normalizar_columnas(df.copy())
    faltantes = [c for c in REQUERIDAS if c not in df.columns]
    if faltantes:
        raise ValueError(f"No encontré estas columnas: {', '.join(faltantes)}. Columnas recibidas: {list(df.columns)}")
    stats = {"filas_entrada": len(df)}
    df["fila_original"] = range(2, len(df) + 2)   # número de fila como se ve en Excel/CSV
    df["fecha"] = df["fecha"].map(lambda v: parsear_fecha(v, dia_primero))
    df["monto"] = df["monto"].map(parsear_monto)
    df["cliente"] = df["cliente"].map(lambda v: limpiar_texto(v, titulo=True))
    df["producto"] = df["producto"].map(limpiar_texto)
    df["cliente"], cambios_cliente = unificar_variantes(df["cliente"])
    df["producto"], cambios_producto = unificar_variantes(df["producto"])
    stats["nombres_unificados"] = cambios_cliente + cambios_producto
    if "cantidad" in df.columns:
        df["cantidad"] = df["cantidad"].map(parsear_monto)
    else:
        df["cantidad"] = 1.0

    motivos = pd.Series("", index=df.index)
    for col, msg in [("fecha", "fecha vacía o ilegible"), ("cliente", "cliente vacío"),
                     ("producto", "producto vacío"), ("monto", "monto vacío o ilegible")]:
        motivos[df[col].isna()] += msg + "; "
    motivos[(df["cantidad"].isna()) | (df["cantidad"] <= 0)] += "cantidad inválida; "
    problemas = df[motivos != ""].assign(motivo=motivos[motivos != ""].str.rstrip("; "))
    limpios = df[motivos == ""].copy()

    columnas = ["fecha", "cliente", "producto", "cantidad", "monto"]
    antes = len(limpios)
    duplicados = limpios[limpios.duplicated(subset=columnas, keep="first")]
    limpios = limpios.drop_duplicates(subset=columnas, keep="first")
    if len(duplicados):
        problemas = pd.concat([problemas, duplicados.assign(motivo="duplicado exacto (se eliminó)")])
    stats.update(filas_limpias=len(limpios), filas_con_problemas=int((motivos != "").sum()),
                 duplicados_eliminados=antes - len(limpios))
    limpios = limpios.sort_values(["fecha", "cliente"]).reset_index(drop=True)
    limpios["mes"] = limpios["fecha"].map(lambda d: date(d.year, d.month, 1))
    return limpios[["fecha", "mes", "cliente", "producto", "cantidad", "monto", "fila_original"]], \
        problemas.sort_values("fila_original")[["fila_original", "motivo"] + [c for c in columnas if c in problemas]], stats


def escribir_reporte(limpios: pd.DataFrame, problemas: pd.DataFrame, stats: dict, salida: Path) -> None:
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    mensual = limpios.groupby("mes", as_index=False).agg(ventas=("monto", "size"), unidades=("cantidad", "sum"),
                                                        monto=("monto", "sum"))
    por_cliente = limpios.groupby("cliente", as_index=False).agg(ventas=("monto", "size"), monto=("monto", "sum")) \
        .sort_values("monto", ascending=False)
    por_producto = limpios.groupby("producto", as_index=False).agg(unidades=("cantidad", "sum"), monto=("monto", "sum")) \
        .sort_values("monto", ascending=False)
    resumen = pd.DataFrame({"concepto": ["Filas recibidas", "Filas limpias", "Filas con problemas",
                                         "Duplicados eliminados", "Nombres unificados", "Monto total limpio"],
                            "valor": [stats["filas_entrada"], stats["filas_limpias"], stats["filas_con_problemas"],
                                      stats["duplicados_eliminados"], stats["nombres_unificados"],
                                      round(float(limpios["monto"].sum()), 2)]})
    hojas = {"Resumen": resumen, "Datos limpios": limpios, "Por mes": mensual,
             "Por cliente": por_cliente, "Por producto": por_producto, "Problemas": problemas}
    with pd.ExcelWriter(salida, engine="openpyxl") as xw:
        for nombre, tabla in hojas.items():
            tabla.to_excel(xw, sheet_name=nombre, index=False)
        wb = xw.book
        for ws in wb.worksheets:
            for celda in ws[1]:
                celda.font = Font(name="Arial", bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F3A5F")
                celda.alignment = Alignment(horizontal="center")
            for fila in ws.iter_rows(min_row=2):
                for celda in fila:
                    celda.font = Font(name="Arial")
                    titulo = ws.cell(row=1, column=celda.column).value
                    if titulo in ("monto", "valor") and isinstance(celda.value, float):
                        celda.number_format = '"$"#,##0.00'
                    elif titulo == "fecha":
                        celda.number_format = "yyyy-mm-dd"
                    elif titulo == "mes":
                        celda.number_format = "mmm yyyy"
            for i, col in enumerate(ws.columns, 1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 45)
            ws.freeze_panes = "A2"
        ws = wb["Por mes"]
        grafica = BarChart()
        grafica.title, grafica.legend, grafica.height, grafica.width = "Ventas por mes", None, 7.5, 16
        grafica.add_data(Reference(ws, min_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
        grafica.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(grafica, "F2")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Limpia una exportación de ventas y genera un reporte en Excel.")
    ap.add_argument("entrada", type=Path, help="Archivo CSV o Excel de entrada")
    ap.add_argument("--salida", type=Path, default=Path("reporte_limpio.xlsx"))
    ap.add_argument("--mes-primero", action="store_true", help="Leer fechas ambiguas como mes/día/año (EE. UU.)")
    args = ap.parse_args(argv)

    if args.entrada.suffix.lower() in (".xlsx", ".xls"):
        crudo = pd.read_excel(args.entrada, dtype=object)
    else:
        crudo = pd.read_csv(args.entrada, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        crudo = crudo.replace({"": None})
    limpios, problemas, stats = limpiar(crudo, dia_primero=not args.mes_primero)
    escribir_reporte(limpios, problemas, stats, args.salida)
    print(f"Filas recibidas: {stats['filas_entrada']} | limpias: {stats['filas_limpias']} | "
          f"con problemas: {stats['filas_con_problemas']} | duplicados eliminados: {stats['duplicados_eliminados']} | "
          f"nombres unificados: {stats['nombres_unificados']}")
    print(f"Reporte: {args.salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
